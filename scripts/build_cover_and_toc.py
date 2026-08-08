from pathlib import Path
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\ADMIN\OneDrive\Tài liệu\CNPM 33")
OUT_DIR = ROOT / "output" / "documents"
TMP_DIR = ROOT / "tmp" / "report_frontmatter"
OUT_PATH = OUT_DIR / "Bao_cao_KhIm_Hub_Bia_va_Muc_luc.docx"
SOURCE_IMAGE = Path(r"C:\Users\ADMIN\AppData\Local\Temp\codex-clipboard-a6c1855b-3689-47dc-bb3b-f81c418b21e5.png")
LOGO_PATH = TMP_DIR / "logo_ntt_niie.png"


def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, size=13, bold=False, italic=False, color="000000"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_centered(doc, text="", size=13, bold=False, after=0, before=0, line=1.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold)
    return p


def add_toc_line(doc, text, level=0, major=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm({0: 0, 1: 0.65, 2: 1.30}.get(level, 1.30))
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(3 if major else 0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_with_next = major
    usable = Cm(16)
    p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(text)
    set_run_font(r, size=11.5 if level == 0 else 11, bold=major)
    p.add_run("\t")
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=10)


def crop_logo():
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    with Image.open(SOURCE_IMAGE) as im:
        # The source is the supplied cover screenshot (376 x 537). Crop both marks.
        crop = im.crop((115, 98, 273, 190))
        crop.save(LOGO_PATH)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    for style_name, size, before, after in (
        ("Heading 1", 16, 12, 8),
        ("Heading 2", 13, 10, 5),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_cover(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    add_centered(doc, "TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH", 13, True, after=5)
    add_centered(doc, "VIỆN ĐÀO TẠO QUỐC TẾ NTT (NIIE)", 13, True, after=5)
    add_centered(doc, "---o0o---", 11, False, after=6)
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(8)
    p_logo.add_run().add_picture(str(LOGO_PATH), width=Cm(5.3))

    add_centered(doc, "ĐỒ ÁN MÔN HỌC", 14, True, after=8)
    add_centered(doc, "CHUYÊN ĐỀ CÔNG NGHỆ PHẦN MỀM 3", 13, True, after=14)
    add_centered(doc, "XÂY DỰNG NỀN TẢNG QUẢN LÝ VÀ CHIA SẺ", 15, True, after=3)
    add_centered(doc, "HỌC LIỆU SỐ KHIM HUB", 16, True, after=3)
    add_centered(doc, "SỬ DỤNG NODE.JS VÀ REACT", 15, True, after=14)

    add_centered(doc, "Giảng viên hướng dẫn: ........................................................", 12.5, after=5)
    add_centered(doc, "Sinh viên thực hiện", 12.5, True, after=6)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(8.4), Cm(5.2)]
    for row in table.rows:
        row.height = Cm(0.85)
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    labels = [("Họ và tên", "Mã số sinh viên"), ("................................", "................................")]
    for ridx, values in enumerate(labels):
        for cidx, value in enumerate(values):
            p = table.cell(ridx, cidx).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=11.5, bold=(ridx == 0))
            if ridx == 0:
                set_cell_shading(table.cell(ridx, cidx), "F2F2F2")

    add_centered(doc, "TP. HỒ CHÍ MINH, tháng 8 năm 2026", 11.5, False, before=58)


TOC = [
    ("LỜI CAM ĐOAN", 0, True), ("LỜI CẢM ƠN", 0, True),
    ("NHẬN XÉT CỦA GIẢNG VIÊN", 0, True), ("MỤC LỤC", 0, True),
    ("DANH MỤC TỪ VIẾT TẮT", 0, True), ("DANH MỤC HÌNH ẢNH", 0, True),
    ("DANH MỤC BẢNG BIỂU", 0, True), ("LỜI MỞ ĐẦU", 0, True),
    ("1. Lý do chọn đề tài", 1, False), ("2. Mục tiêu đề tài", 1, False),
    ("3. Đối tượng nghiên cứu", 1, False), ("4. Phạm vi nghiên cứu", 1, False),
    ("5. Phương pháp thực hiện", 1, False), ("6. Kết cấu báo cáo", 1, False),
    ("CHƯƠNG 1: TỔNG QUAN DỰ ÁN VÀ CƠ SỞ LÝ THUYẾT", 0, True),
    ("1.1. Bối cảnh dự án", 1, False), ("1.1.1. Thực trạng lưu trữ học liệu số", 2, False),
    ("1.1.2. Nhu cầu chia sẻ và khai thác học liệu", 2, False), ("1.1.3. Các vấn đề cần giải quyết", 2, False),
    ("1.2. Giới thiệu đề tài KhIm Hub", 1, False), ("1.2.1. Mục tiêu hệ thống", 2, False),
    ("1.2.2. Đối tượng sử dụng", 2, False), ("1.2.3. Phạm vi chức năng", 2, False),
    ("1.2.4. Ý nghĩa thực tiễn", 2, False),
    ("1.3. Khảo sát các nền tảng học liệu số tương tự", 1, False),
    ("1.3.1. Moodle", 2, False), ("1.3.2. Google Classroom", 2, False),
    ("1.3.3. Thư viện số của trường đại học", 2, False), ("1.3.4. So sánh và đánh giá", 2, False),
    ("1.4. Công nghệ Backend", 1, False), ("1.4.1. Node.js", 2, False),
    ("1.4.2. Express.js", 2, False), ("1.4.3. RESTful API", 2, False),
    ("1.5. Công nghệ Frontend", 1, False), ("1.5.1. ReactJS", 2, False),
    ("1.5.2. Vite", 2, False), ("1.5.3. React Router", 2, False),
    ("1.5.4. Axios và thư viện giao diện", 2, False),
    ("1.6. Cơ sở dữ liệu và lưu trữ tệp", 1, False), ("1.6.1. MySQL", 2, False),
    ("1.6.2. Thiết kế dữ liệu quan hệ", 2, False), ("1.6.3. Lưu trữ tệp học liệu", 2, False),
    ("1.7. Xác thực và bảo mật", 1, False), ("1.7.1. JSON Web Token (JWT)", 2, False),
    ("1.7.2. Mã hóa mật khẩu", 2, False), ("1.7.3. Phân quyền theo vai trò", 2, False),
    ("1.8. Công cụ phát triển", 1, False), ("1.8.1. Visual Studio Code", 2, False),
    ("1.8.2. Git và GitHub", 2, False), ("1.8.3. Postman", 2, False),
    ("1.9. Lý do lựa chọn công nghệ", 1, False), ("1.10. Kết chương", 1, False),
    ("CHƯƠNG 2: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", 0, True),
    ("2.1. Phân tích yêu cầu hệ thống", 1, False), ("2.1.1. Yêu cầu của sinh viên", 2, False),
    ("2.1.2. Yêu cầu của giảng viên", 2, False), ("2.1.3. Yêu cầu của quản trị viên", 2, False),
    ("2.1.4. Yêu cầu phi chức năng", 2, False),
    ("2.2. Phân tích Use Case", 1, False), ("2.2.1. Biểu đồ Use Case tổng quát", 2, False),
    ("2.2.2. Đặc tả Use Case đăng nhập", 2, False), ("2.2.3. Đặc tả Use Case tải lên học liệu", 2, False),
    ("2.2.4. Đặc tả Use Case duyệt học liệu", 2, False), ("2.2.5. Đặc tả Use Case tìm kiếm và lọc", 2, False),
    ("2.2.6. Đặc tả Use Case xem và tải tài liệu", 2, False),
    ("2.3. Thiết kế kiến trúc hệ thống", 1, False), ("2.3.1. Kiến trúc tổng thể", 2, False),
    ("2.3.2. Kiến trúc Client–Server", 2, False), ("2.3.3. Luồng xử lý dữ liệu", 2, False),
    ("2.4. Thiết kế cơ sở dữ liệu", 1, False), ("2.4.1. Biểu đồ ERD", 2, False),
    ("2.4.2. Data Dictionary", 2, False), ("2.4.3. Quan hệ giữa các bảng", 2, False),
    ("2.5. Thiết kế biểu đồ tuần tự", 1, False), ("2.5.1. Luồng đăng nhập", 2, False),
    ("2.5.2. Luồng tải lên và duyệt học liệu", 2, False), ("2.5.3. Luồng tìm kiếm và tải tài liệu", 2, False),
    ("2.6. Thiết kế giao diện người dùng", 1, False), ("2.6.1. Giao diện công khai", 2, False),
    ("2.6.2. Giao diện sinh viên", 2, False), ("2.6.3. Giao diện giảng viên", 2, False),
    ("2.6.4. Giao diện quản trị viên", 2, False), ("2.7. Thiết kế bảo mật", 1, False),
    ("2.8. Kết chương", 1, False),
    ("CHƯƠNG 3: TRIỂN KHAI ỨNG DỤNG VÀ KIỂM THỬ", 0, True),
    ("3.1. Môi trường phát triển và cấu trúc mã nguồn", 1, False),
    ("3.1.1. Công cụ và môi trường", 2, False), ("3.1.2. Cấu trúc thư mục Backend", 2, False),
    ("3.1.3. Cấu trúc thư mục Frontend", 2, False),
    ("3.2. Triển khai cơ sở dữ liệu MySQL", 1, False), ("3.2.1. Khởi tạo cơ sở dữ liệu", 2, False),
    ("3.2.2. Tạo bảng và dữ liệu mẫu", 2, False),
    ("3.3. Phân hệ xác thực và phân quyền", 1, False), ("3.3.1. Đăng nhập", 2, False),
    ("3.3.2. Xác thực JWT", 2, False), ("3.3.3. Phân quyền sinh viên, giảng viên và quản trị viên", 2, False),
    ("3.4. Phân hệ quản lý học liệu", 1, False), ("3.4.1. Quản lý danh mục và môn học", 2, False),
    ("3.4.2. Tải lên tệp PDF, video và slide", 2, False), ("3.4.3. Kiểm duyệt nội dung", 2, False),
    ("3.4.4. Quản lý trạng thái học liệu", 2, False),
    ("3.5. Phân hệ tìm kiếm và lọc", 1, False), ("3.5.1. Tìm kiếm theo từ khóa", 2, False),
    ("3.5.2. Lọc theo môn học và loại tài liệu", 2, False),
    ("3.6. Phân hệ người dùng", 1, False), ("3.6.1. Xem chi tiết học liệu", 2, False),
    ("3.6.2. Tải tài liệu", 2, False), ("3.6.3. Theo dõi lịch sử xem và tải", 2, False),
    ("3.7. Phân hệ quản trị", 1, False), ("3.7.1. Thống kê tổng quan", 2, False),
    ("3.7.2. Duyệt hoặc từ chối học liệu", 2, False), ("3.7.3. Quản lý người dùng", 2, False),
    ("3.8. Triển khai giao diện React", 1, False), ("3.8.1. Trang chủ KhIm Hub", 2, False),
    ("3.8.2. Trang đăng nhập", 2, False), ("3.8.3. Dashboard theo vai trò", 2, False),
    ("3.8.4. Thiết kế responsive", 2, False),
    ("3.9. Kỹ thuật nâng cao đã áp dụng", 1, False), ("3.9.1. Kiểm soát kích thước và định dạng tệp", 2, False),
    ("3.9.2. Bảo vệ API bằng JWT", 2, False), ("3.9.3. Ghi nhận lượt xem và lượt tải", 2, False),
    ("3.10. Kiểm thử phần mềm", 1, False), ("3.10.1. Kế hoạch và phương pháp kiểm thử", 2, False),
    ("3.10.2. Các kịch bản kiểm thử tiêu biểu", 2, False), ("3.10.3. Kết quả kiểm thử", 2, False),
    ("3.11. Kịch bản trình diễn hệ thống", 1, False), ("3.12. Kết chương", 1, False),
    ("CHƯƠNG 4: TỔNG KẾT VÀ BÀI HỌC KINH NGHIỆM", 0, True),
    ("4.1. Kết quả đạt được", 1, False), ("4.2. So sánh kết quả với mục tiêu ban đầu", 1, False),
    ("4.3. Khó khăn trong quá trình thực hiện", 1, False), ("4.4. Hạn chế của hệ thống", 1, False),
    ("4.5. Bài học kinh nghiệm", 1, False), ("4.6. Hướng phát triển", 1, False),
    ("4.6.1. Tích hợp AI gợi ý học liệu", 2, False), ("4.6.2. Thảo luận và đánh giá học liệu", 2, False),
    ("4.6.3. Lưu trữ đám mây và streaming video", 2, False), ("4.7. Kết luận", 1, False),
    ("TÀI LIỆU THAM KHẢO", 0, True), ("PHỤ LỤC", 0, True),
    ("Phụ lục A. Link mã nguồn GitHub/GitLab", 1, False),
    ("Phụ lục B. Hướng dẫn cài đặt và cấu hình môi trường", 1, False),
    ("Phụ lục C. Cấu trúc cơ sở dữ liệu và dữ liệu mẫu", 1, False),
    ("Phụ lục D. Danh sách API", 1, False), ("Phụ lục E. Bộ test case", 1, False),
    ("Phụ lục F. Hình ảnh giao diện hệ thống", 1, False),
]


def build_toc(doc):
    doc.add_page_break()
    section = doc.sections[0]
    # First-page footer is suppressed; later pages use centered page numbers.
    add_page_number(section.footer.paragraphs[0])
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("MỤC LỤC"), size=16, bold=True)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(10)
    set_run_font(note.add_run("(Số trang được cập nhật sau khi hoàn thiện nội dung báo cáo)"), size=10.5, italic=True, color="666666")
    for text, level, major in TOC:
        add_toc_line(doc, text, level, major)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    crop_logo()
    doc = Document()
    configure_styles(doc)
    build_cover(doc)
    build_toc(doc)
    props = doc.core_properties
    props.title = "Bìa và mục lục báo cáo KhIm Hub"
    props.subject = "Đồ án môn học Chuyên đề Công nghệ Phần mềm 3"
    props.author = "Sinh viên thực hiện"
    doc.save(OUT_PATH)
    print("DOCX created successfully")


if __name__ == "__main__":
    main()
